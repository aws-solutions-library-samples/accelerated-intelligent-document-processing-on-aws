# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0

"""A legacy `.xls` upload must reach the redactor as something it can read.

The hook claims `.xls` (the UI accepts it, and the dispatch routes it), but the
vendored redactor reads workbooks with `openpyxl`, which handles only the OOXML
container — on a BIFF file it fails with `zipfile.BadZipFile`. Because the
extension is *claimed*, it never reaches the handler's unsupported-format path
either: it enters the xlsx branch and raises, and the shipped preset's
`onError: fail` turns that into a failed run. `handler._convert_xls_to_xlsx`
rewrites the workbook first (#824).

These tests exercise the real converter against a real BIFF file — an import
assertion would not catch this class of bug, since the failure is openpyxl
refusing the container. S3 is faked at the boundary (`get_object`/`put_object`)
rather than mocked with moto, matching the rest of this suite, which never needs
AWS.

The fixture is the same synthetic two-sheet workbook used by the host's Excel
converter test (`lib/idp_common_pkg/tests/unit/fixtures/two_sheets.xls`, added in
GitHub #800, generated with `xlwt`). It is duplicated here on purpose: this
function is built, linted and tested as its own package, with its own
requirements, so its tests should not reach across the repo.
"""

import importlib
import io
import os
import sys

import pytest

HOOK_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
FIXTURE = os.path.join(os.path.dirname(__file__), "fixtures", "two_sheets.xls")


@pytest.fixture(autouse=True)
def _env(monkeypatch):
    monkeypatch.setenv("INPUT_BUCKET", "input-bkt")
    monkeypatch.setenv("WORKING_BUCKET", "working-bkt")
    monkeypatch.setenv("REDACTED_SUFFIX", "(REDACTED)")
    sys.path.insert(0, HOOK_DIR)
    yield
    sys.path.remove(HOOK_DIR)
    sys.modules.pop("handler", None)


def _load():
    if "handler" in sys.modules:
        del sys.modules["handler"]
    return importlib.import_module("handler")


class _FakeS3:
    """Just the two calls the converter makes."""

    def __init__(self, objects: dict[tuple[str, str], bytes]) -> None:
        self.objects = objects
        self.puts: dict[tuple[str, str], bytes] = {}

    def get_object(self, Bucket: str, Key: str):  # noqa: N803 - boto3 kwarg names
        return {"Body": io.BytesIO(self.objects[(Bucket, Key)])}

    def put_object(self, Bucket: str, Key: str, Body: bytes):  # noqa: N803
        self.puts[(Bucket, Key)] = Body
        return {}


def _converted_workbook(monkeypatch, mod, raw: bytes):
    """Run the converter over ``raw`` and return the openpyxl workbook it wrote."""
    from openpyxl import load_workbook

    fake = _FakeS3({("input-bkt", "src/legacy.xls"): raw})
    monkeypatch.setattr(mod, "_s3", fake)

    key = mod._convert_xls_to_xlsx(
        "input-bkt", "src/legacy.xls", "pii_scratch/doc-1/legacy.converted.xlsx"
    )

    assert key == "pii_scratch/doc-1/legacy.converted.xlsx"
    assert ("working-bkt", key) in fake.puts, (
        "conversion must land in the Working bucket"
    )
    return load_workbook(io.BytesIO(fake.puts[("working-bkt", key)]))


def test_the_conversion_is_readable_by_the_engine_the_redactor_uses(monkeypatch):
    """The whole point: openpyxl can open the result. On the unconverted file the
    same call fails (see the test below)."""
    mod = _load()
    with open(FIXTURE, "rb") as handle:
        raw = handle.read()

    workbook = _converted_workbook(monkeypatch, mod, raw)

    assert len(workbook.worksheets) == 2
    text = " ".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    # Content from BOTH sheets survives — a converter that silently dropped
    # sheets would still produce a loadable workbook.
    assert "Widget" in text
    assert "GrandTotal" in text


def test_openpyxl_cannot_read_the_original():
    """Pins the premise rather than trusting it: without conversion the vendored
    processor's engine refuses the file, which is why `.xls` failed the run.

    It raises ``BadZipFile``, not the ``InvalidFileException`` you might expect:
    that check is on the *filename*, and the vendored processor downloads to a
    temp path it names ``..._<base>.xlsx`` (`tabular_processor.py`), so the
    extension check passes and openpyxl gets as far as unzipping a BIFF file.
    """
    from zipfile import BadZipFile

    from openpyxl import load_workbook

    with open(FIXTURE, "rb") as handle:
        raw = handle.read()

    with pytest.raises(BadZipFile):
        load_workbook(io.BytesIO(raw))


def test_sheet_names_are_preserved(monkeypatch):
    mod = _load()
    with open(FIXTURE, "rb") as handle:
        workbook = _converted_workbook(monkeypatch, mod, handle.read())

    import xlrd

    with open(FIXTURE, "rb") as handle:
        original = xlrd.open_workbook(file_contents=handle.read())

    assert workbook.sheetnames == [sheet.name for sheet in original.sheets()]


def test_cell_values_keep_their_type():
    """Excel stores every number as a float; writing 42.0 where the sheet showed
    42 changes what the detector scans and what a reviewer reads."""
    mod = _load()
    xlrd = pytest.importorskip("xlrd")

    class _Cell:
        def __init__(self, ctype, value):
            self.ctype, self.value = ctype, value

    datemode = 0
    assert mod._xls_cell_value(_Cell(xlrd.XL_CELL_NUMBER, 42.0), datemode) == 42
    assert isinstance(
        mod._xls_cell_value(_Cell(xlrd.XL_CELL_NUMBER, 42.0), datemode), int
    )
    assert mod._xls_cell_value(_Cell(xlrd.XL_CELL_NUMBER, 42.5), datemode) == 42.5
    assert (
        mod._xls_cell_value(_Cell(xlrd.XL_CELL_TEXT, "Jane Doe"), datemode)
        == "Jane Doe"
    )
    assert mod._xls_cell_value(_Cell(xlrd.XL_CELL_BOOLEAN, 1), datemode) is True
    assert mod._xls_cell_value(_Cell(xlrd.XL_CELL_EMPTY, ""), datemode) is None
    assert mod._xls_cell_value(_Cell(xlrd.XL_CELL_ERROR, 42), datemode) is None
    # A date becomes a datetime, which openpyxl writes as a real date cell.
    converted = mod._xls_cell_value(_Cell(xlrd.XL_CELL_DATE, 45000.0), datemode)
    assert converted.year == 2023
    # An out-of-range serial must not fail the redaction; the raw number is kept.
    assert mod._xls_cell_value(_Cell(xlrd.XL_CELL_DATE, 1e9), datemode) == 1e9


def _stub_processor(monkeypatch, seen):
    """Replace the vendored Excel processor with one that records where it was
    pointed and returns the result shape ``_redact_to_scratch`` requires."""

    def _fake_process_excel_file(source_bucket, source_key, *_args, **_kwargs):
        seen["source"] = (source_bucket, source_key)
        return {
            "success": True,
            "s3_output_file": "pii_scratch/doc/redacted_sheet.xlsx",
            "pii_count": 2,
        }

    module = type(sys)("processors.tabular_processor")
    module.process_excel_file = _fake_process_excel_file
    monkeypatch.setitem(sys.modules, "processors.tabular_processor", module)


def test_the_xls_branch_converts_before_calling_the_processor(monkeypatch):
    """Wiring: the processor must be handed the CONVERTED object in the Working
    bucket, and the original must be left alone — the halt/delete decision
    downstream still refers to it."""
    mod = _load()
    seen: dict = {}

    def _fake_convert(bucket, key, dest):
        seen["convert"] = (bucket, key, dest)
        return dest

    monkeypatch.setattr(mod, "_convert_xls_to_xlsx", _fake_convert)
    _stub_processor(monkeypatch, seen)

    out = mod._redact_to_scratch(
        {"input_bucket": "input-bkt", "input_key": "in/legacy.xls"},
        mod._build_pii_config({}),
        "doc-1",
    )

    assert seen["convert"] == (
        "input-bkt",
        "in/legacy.xls",
        "pii_scratch/doc-1/legacy.converted.xlsx",
    )
    assert seen["source"] == ("working-bkt", "pii_scratch/doc-1/legacy.converted.xlsx")
    # And the redacted copy is still an .xlsx, as it was for both extensions before.
    assert out["out_ext"] == "xlsx"


def test_an_xlsx_is_not_converted(monkeypatch):
    """The common case must not gain a conversion round-trip."""
    mod = _load()
    seen: dict = {}
    calls: list = []

    monkeypatch.setattr(
        mod, "_convert_xls_to_xlsx", lambda *a: calls.append(a) or "unused"
    )
    _stub_processor(monkeypatch, seen)

    mod._redact_to_scratch(
        {"input_bucket": "input-bkt", "input_key": "in/modern.xlsx"},
        mod._build_pii_config({}),
        "doc-2",
    )

    assert calls == []
    assert seen["source"] == ("input-bkt", "in/modern.xlsx")
